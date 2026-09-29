"""Build the Supporting Information (SI.docx) from analysis/final and analysis/figures/out (read-only on both).

Word fields are used for captions (SEQ Figure / SEQ Table) and for the lists of figures and tables (TOC \\c),
so page numbers are computed by Word when the document is opened and exported (see export_si.ps1).
Usage (from the project root): python analysis/si/build_si.py
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FINAL, FIGS, OUT = ROOT / "analysis" / "final", ROOT / "analysis" / "figures" / "out", ROOT / "analysis" / "si"
TITLE = ("Supporting Information for: Validation Noise, Not the Optimizer, Limits Metaheuristic Tuning in "
         "Molecular Property Prediction")
AUTHOR = "Hussein Fadhil Abedi"
DS = ["ESOL", "FreeSolv", "Lipophilicity", "BBBP", "BACE"]
OPTS = ["Random", "TPE", "GWO", "WOA", "PSO"]
METHODS = ["Default"] + OPTS
FS_NAMES = {"random": "Random", "tpe": "TPE", "gwo_ref": "GWO", "woa_ref": "WOA", "pso": "PSO"}
MARGIN = Inches(0.75)
FIG_MAX_W, FIG_MAX_H = 7.0, 7.7  # inches (portrait text area is 7.0 x 9.5 in; leaves room for caption + legend)


# ------------------------------------------------------------------------------------------------ formatting
def f3(x) -> str:
    """Three significant digits, fixed notation where sensible."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return "–"
    x = float(x)
    if x == 0:
        return "0"
    a = abs(x)
    if a >= 100:
        return f"{x:.0f}"
    digits = max(0, 2 - math.floor(math.log10(a)))
    return f"{x:.{digits}f}" if digits <= 8 else f"{x:.2e}"  # fixed notation down to 1e-6


def fp(p) -> str:
    return "–" if p is None or (isinstance(p, float) and math.isnan(p)) else ("<0.001" if p < 1e-3 else f3(p))


def ci(m, lo, hi) -> str:
    return f"{f3(m)} [{f3(lo)}, {f3(hi)}]"


# ------------------------------------------------------------------------------------------------ docx helpers
def add_field(paragraph, instr: str, placeholder: str = "") -> None:
    """Complex field (begin / instr / separate / result / end) so Word can update it."""
    def run_with(el):
        r = OxmlElement("w:r")
        r.append(el)
        paragraph._p.append(r)
        return r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); run_with(b)
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = f" {instr} "; run_with(t)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); run_with(s)
    rt = OxmlElement("w:t"); rt.text = placeholder; run_with(rt)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); run_with(e)


def caption(doc, label: str, number: int, title: str):
    """'Figure S1. Title.' / 'Table S1. Title.' in Caption style with a SEQ field (collected by the TOC)."""
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(f"{label} S"); r.bold = True
    add_field(p, f"SEQ {label} \\* ARABIC", str(number))
    for run in p.runs:
        run.bold = True
    r2 = p.add_run(". "); r2.bold = True
    p.add_run(title)
    p.paragraph_format.keep_with_next = True
    return p


def legend(doc, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.size = Pt(10)
    return p


def new_section(doc, landscape: bool):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Inches(11), Inches(8.5)) if landscape else (Inches(8.5), Inches(11))
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, MARGIN)
    return s


def column_widths(header: list[str], rows: list[list[str]], total_in: float, font_pt: float) -> list[float]:
    """Every column is at least as wide as its longest unbreakable word (so no word is split); the remaining width
    goes to columns whose full entries are longer than that minimum."""
    char_in, pad = 0.52 * font_pt / 72, 0.14          # average glyph width of Times at font_pt, cell padding
    minw, want = [], []
    for j, h in enumerate(header):
        cells = [str(r[j]) for r in rows]
        longest_word = max([len(w) for c in cells for w in c.split()] + [len(w) for w in h.split()] + [2])
        longest_cell = max([len(c) for c in cells] + [longest_word])
        minw.append(longest_word * char_in + pad)
        want.append(min(longest_cell, 42) * char_in + pad)
    extra = total_in - sum(minw)
    if extra <= 0:
        return [total_in * m / sum(minw) for m in minw]
    grow = [w - m for w, m in zip(want, minw)]
    if sum(grow) <= extra:                            # everything fits on one line: spread the slack evenly
        return [w + (extra - sum(grow)) / len(want) for w in want]
    return [m + extra * g / sum(grow) for m, g in zip(minw, grow)]


def add_table(doc, header: list[str], rows: list[list[str]], font_pt: float = 8.0, total_in: float = 7.0):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, h in zip(t.rows[0].cells, header):
        cell.text = ""
        run = cell.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(font_pt)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader"); hdr.set(qn("w:val"), "true"); tr_pr.append(hdr)  # repeat header row
    for row in rows:
        cells = t.add_row().cells
        for cell, v in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(v)); run.font.size = Pt(font_pt)
    widths = column_widths(header, rows, total_in, font_pt)
    for j, col in enumerate(t.columns):
        col.width = Inches(widths[j])
    for row in t.rows:
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            pf = cell.paragraphs[0].paragraph_format
            pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.0
    return t


def order_key(df: pd.DataFrame, cols=("dataset", "experiment", "method")) -> pd.DataFrame:
    df = df.copy()
    if "dataset" in df:
        df["_d"] = df.dataset.map({d: i for i, d in enumerate(DS)})
    if "method" in df:
        df["_m"] = df.method.map({m: i for i, m in enumerate(METHODS)})
    keys = [k for k in ("_d", "experiment", "_r", "metric", "_m") if k in df]
    return df.sort_values(keys).drop(columns=[c for c in ("_d", "_m", "_r") if c in df])


# ------------------------------------------------------------------------------------------------ content
def check_convergence_claim() -> None:
    """Assert the Figure S5 caption claim from the raw call logs: median best-so-far < default by call 10."""
    sys.path.insert(0, str(ROOT / "analysis" / "figures"))
    from data import load_curves, load_study  # read-only
    study = load_study()
    curves = load_curves(study)
    vr = pd.DataFrame(json.loads((FINAL / "final_stats.json").read_text(encoding="utf-8"))["validation_reliability"])
    for ds in DS:
        dval = vr[(vr.dataset == ds) & (vr.experiment == "A") & (vr.method == "Default")].median_val.iloc[0]
        base = (1 - dval) if ds in ("BBBP", "BACE") else dval
        for exp in ("A", "B"):
            default_fit = base + (0.01 if exp == "B" else 0.0)
            for m in OPTS:
                assert np.median(curves[(ds, exp, m)][:, 9]) < default_fit, (ds, exp, m)


FIGURES = [
    ("figS1_sphere.png", "Sphere-function verification of the optimizer implementations.",
     "Median best fitness after 100 evaluations over 30 seeds (lower is better). (a, b) Optimum at 0.5 in every "
     "dimension (filled markers) versus a shifted optimum at a fixed random point in [0.1, 0.9]^d (open markers), for "
     "d = 10 and d = 200; bars show the interquartile range. (c, d) Optimum at 0.2 (filled) versus 0.8 (open) in every "
     "dimension. Faded markers: the a0 variants of GWO and WOA (sensitivity analysis only). At d = 200 with the shifted "
     "optimum, PSO (14.4) and GWO (14.7) reach the lowest values and Random the highest (21.4). GWO does worse when the "
     "optimum is at 0.8 than at 0.2 (ratio 1.60 at d = 200; 1.46 for the a0 variant), consistent with a bias toward "
     "x = 0, whereas Random, TPE and PSO show ratios of 0.986–1.01. WOA reaches unusually low values when the optimum "
     "lies on the diagonal (2.84 and 2.22 at d = 200)."),
    ("figS2_a12_heatmap.png", "Pairwise Vargha–Delaney A12 between optimizers on the primary test metric.",
     "Lower triangle: A12 of the row optimizer versus the column optimizer, that is, the probability that the row "
     "optimizer scores better on the test set (0.5 = no difference), from 30 seeds per method, for each dataset "
     "(columns) and experiment (rows). Boxed cells with an asterisk are significant after Holm correction over the 10 "
     "pairs (paired Wilcoxon signed-rank test, zero_method = pratt). Only 2 of the 100 comparisons are significant, "
     "both for FreeSolv in Experiment A: GWO outperforms Random (cell GWO vs Random, A12 = 0.67) and WOA (cell WOA vs "
     "GWO, A12 = 0.30, i.e., 0.70 in favour of GWO). All values are listed in Table S3."),
    ("figS3_cd_optimizers.png", "Critical-difference diagrams across datasets for the five optimizers (low-powered).",
     "Mean rank of each optimizer across the N = 5 datasets (rank 1 = best median test score on the primary metric), "
     "with the Nemenyi critical difference at α = 0.05 (CD = 2.73); thick bars join optimizers whose mean ranks differ "
     "by less than the CD. (a) Experiment A: GWO 2.00, WOA 2.40, PSO 2.60, TPE 3.40, Random 4.60; Friedman p = 0.075; "
     "no pair differs by more than the CD. (b) Experiment B: PSO 1.60, GWO 2.40, TPE 2.60, Random 3.80, WOA 4.60; "
     "Friedman p = 0.023; only PSO and WOA differ by more than the CD (3.00). With five datasets these tests have low "
     "power; the per-dataset analyses (Tables S2 and S3) are the primary evidence."),
    ("figS4_sensitivity.png", "Sensitivity of GWO and WOA to the coefficient schedule: a0 versus reference-faithful variants.",
     "Paired differences (a0 − reference) in the primary test metric for each of the 30 seeds, oriented so that positive "
     "values favour the a0 schedule; bars mark the median paired difference. Rows are datasets, columns are GWO and "
     "WOA, and each panel shows Experiments A and B. Holm-adjusted p-values (paired Wilcoxon test, Holm correction over "
     "the 10 dataset × experiment cells per method) are shown at right; none is significant (smallest Holm p = 0.417, "
     "GWO, ESOL, Experiment B). The a0 schedule reduces the median number of unique evaluations per run (GWO 91 versus "
     "100; WOA 91.5–93 versus 96–97.5) but changes none of the 50 tuning-versus-default conclusions. Values in Table S7."),
    ("figS5_convergence.png", "Convergence of the best-so-far validation fitness.",
     "Median (line) and interquartile range (band) over 30 seeds of the best-so-far validation fitness against the "
     "fitness-call index (1–100), per dataset (rows) and experiment (columns): RMSE for regression and 1 − ROC-AUC for "
     "classification; in Experiment B the fitness includes the feature-count penalty (0.01 × selected fraction). Dashed "
     "line: validation fitness of the default configuration. The grey vertical line marks the end of the shared "
     "10-point initial design (calls 1–10 are identical for all optimizers). For every optimizer, dataset and "
     "experiment, the median best-so-far fitness is already below the default's by call 10."),
]


def build_tables(fs: dict) -> list[tuple]:
    """(title, header, rows, landscape, note, font_pt)."""
    tables = []
    prev = fs["provenance"].get("pr_auc_test_prevalence", {})

    # S1 tuning vs default
    t = pd.read_csv(FINAL / "1_tuning_vs_default.csv")
    t["_r"] = (t.role != "primary").astype(int)
    t = order_key(t)
    rows = [[r.dataset, r.experiment, f"{r.metric}{' (primary)' if r.role == 'primary' else ''}", r.method, f3(r.default),
             ci(r.median, r.ci_low, r.ci_high), ci(r.median_improvement, r.impr_ci_low, r.impr_ci_high),
             f3(r.pct_seeds_beating_default), fp(r.p_holm), r.conclusion] for r in t.itertuples()]
    tables.append(("Tuning versus the default configuration, all test metrics.",
                   ["Dataset", "Exp.", "Metric", "Method", "Default", "Median [95% CI]", "Improvement [95% CI]",
                    "% seeds better", "p (Holm)", "Conclusion"], rows, True,
                   "Improvement over the default is oriented so that positive values are better (default − method for "
                   "RMSE and MAE; method − default for R², ROC-AUC and PR-AUC). CIs: percentile bootstrap, 10,000 "
                   "resamples of the 30 tuning seeds. p: one-sample Wilcoxon signed-rank test against the default "
                   "value (zero_method = pratt), Holm correction over the 5 optimizers within each dataset × experiment "
                   "× metric. Conclusion: better / worse than the default at Holm p < 0.05, otherwise n.s. PR-AUC chance "
                   f"level (test prevalence): BBBP {f3(prev.get('BBBP', float('nan')))}, BACE {f3(prev.get('BACE', float('nan')))}.",
                   7.0))

    # S2 Friedman
    t = pd.read_csv(FINAL / "2_friedman.csv")
    t["_r"] = (t.role != "primary").astype(int)
    t = order_key(t)
    rows = [[r.dataset, r.experiment, f"{r.metric}{' (primary)' if r.role == 'primary' else ''}", f3(r.friedman_chi2),
             fp(r.friedman_p)] + [f3(getattr(r, f"mean_rank_{m}")) for m in OPTS] for r in t.itertuples()]
    tables.append(("Friedman tests across the five optimizers (default excluded).",
                   ["Dataset", "Exp.", "Metric", "χ²", "p"] + [f"Mean rank {m}" for m in OPTS], rows, True,
                   "Friedman test with the 30 tuning seeds as blocks and the 5 optimizers as treatments (4 degrees of "
                   "freedom); mean rank 1 = best.", 8.0))

    # S3 pairwise (primary)
    t = pd.read_csv(FINAL / "2_pairwise.csv")
    sec = t[t.role == "secondary"]
    t = t[t.role == "primary"].copy()
    t["_d"] = t.dataset.map({d: i for i, d in enumerate(DS)})
    t = t.sort_values(["_d", "experiment"], kind="stable")
    rows = [[r.dataset, r.experiment, r.metric, r.pair, f3(r.median_diff_goodness), fp(r.p_raw), fp(r.p_holm),
             f3(r.A12)] for r in t.itertuples()]
    tables.append(("Pairwise comparisons of optimizers on the primary metric: Wilcoxon signed-rank test with Holm "
                   "correction and Vargha–Delaney A12.",
                   ["Dataset", "Exp.", "Metric", "Pair (first vs second)", "Median paired difference", "p (raw)",
                    "p (Holm)", "A12"], rows, False,
                   "Paired Wilcoxon signed-rank test over 30 seeds (zero_method = pratt), Holm correction over the 10 "
                   "pairs within each dataset × experiment. Median difference and A12 are oriented so that values > 0 "
                   "and > 0.5, respectively, favour the first optimizer of the pair. Secondary metrics (MAE, R², PR-AUC) "
                   f"are in 2_pairwise.csv; {int((sec.p_holm < .05).sum())} of {len(sec)} secondary comparisons are "
                   "Holm-significant.", 8.0))

    # S4 validation reliability
    t = order_key(pd.read_csv(FINAL / "3_validation_reliability.csv"))
    rows = [[r.dataset, r.experiment, r.method, r.metric, f3(r.median_val), f3(r.median_test),
             f3(r.median_gap_test_minus_val), "–" if math.isnan(r.rho) else ci(r.rho, r.rho_ci_low, r.rho_ci_high)]
            for r in t.itertuples()]
    tables.append(("Validation–test gap and Spearman correlation between best validation and test scores.",
                   ["Dataset", "Exp.", "Method", "Metric", "Median best validation", "Median test",
                    "Median gap (test − val)", "Spearman ρ [95% CI]"], rows, False,
                   "Validation score = best validation score of the run (for the default: its validation score), on "
                   "the same scale as the test metric (Experiment B without the feature-count penalty). ρ: Spearman "
                   "correlation between validation and test scores across the 30 seeds, percentile bootstrap CI "
                   "(10,000 resamples); not defined for the deterministic default.", 8.0))

    # S5 feature selection
    t = order_key(pd.read_csv(FINAL / "4_feature_selection.csv"))
    rows = [[r.dataset, int(r.F), r.method, ci(r.fraction_median, r.fraction_ci_low, r.fraction_ci_high),
             f"{f3(r.fraction_q25)}–{f3(r.fraction_q75)}", ci(r.nogueira, r.nogueira_ci_low, r.nogueira_ci_high)]
            for r in t.itertuples()]
    tables.append(("Selected-feature fraction and Nogueira stability (Experiment B).",
                   ["Dataset", "F", "Method", "Selected fraction, median [95% CI]", "IQR",
                    "Nogueira Φ [95% CI]"], rows, False,
                   "F: number of descriptors available after training-split filtering. Selected fraction = selected / F "
                   "per run; CI: percentile bootstrap. Nogueira stability over the 30 selected subsets with its "
                   "asymptotic 95% CI (Nogueira et al., 2018); Φ = 0 corresponds to random selection.", 8.0))

    # S6 mirror
    t = pd.read_csv(FINAL / "5_mirror_prespecified_tests.csv")
    t["method"] = t.optimizer.map(FS_NAMES)
    t["_o"] = t.outcome.map(lambda o: 0 if o == "fraction" else 1)
    t["_d"] = t.dataset.map({d: i for i, d in enumerate(DS)})
    t["_m"] = t.method.map({m: i for i, m in enumerate(OPTS)})
    t = t.sort_values(["_d", "_o", "_m"])
    lab = {"fraction": "Selected fraction", "test_roc_auc": "Test ROC-AUC", "test_rmse": "Test RMSE"}
    rows = [[r.dataset, lab[r.outcome], r.method, f3(r.median_normal), f3(r.median_mirrored), f3(r.median_paired_diff),
             f"{int(r.n_mirrored_higher)}/30", fp(r.p_raw), fp(r.p_holm)] for r in t.itertuples()]
    tables.append(("Mirrored-encoding robustness check (pre-registered tests, Experiment B).",
                   ["Dataset", "Outcome", "Optimizer", "Median normal", "Median mirrored", "Median paired difference",
                    "Seeds mirrored > normal", "p (raw)", "p (Holm)"], rows, True,
                   "Mirrored encoding: every proposed point x is decoded as 1 − x (including the shared initial design). "
                   "Paired Wilcoxon signed-rank test over 30 seeds (zero_method = pratt), difference = mirrored − "
                   "normal, Holm correction over the 5 optimizers within each dataset × outcome, as pre-registered in "
                   "results/xgb_mh/mirror/prediction.md before the run. Decision rule: GWO supports the prediction if its "
                   "mirrored fraction is higher (Holm p < 0.05) and its median mirrored fraction exceeds 0.5 on both "
                   "datasets (met: 4 of 4 conditions); Random (control) shows no change (met).", 8.0))

    # S7 sensitivity
    t = pd.read_csv(FINAL / "6_sensitivity_a0_vs_ref.csv")
    t["_d"] = t.dataset.map({d: i for i, d in enumerate(DS)})
    t = t.sort_values(["method", "_d", "experiment"])
    rows = [[r.method, r.dataset, r.experiment, r.metric, f3(r.median_ref), f3(r.median_a0),
             f3(r.median_paired_diff_ref_better), fp(r.p_raw), fp(r.p_holm),
             f"{f3(r.unique_ref_median)} / {f3(r.unique_a0_median)}"] for r in t.itertuples()]
    rq1 = pd.DataFrame(fs["sensitivity_rq1"])
    tables.append(("Sensitivity analysis: a0 versus reference-faithful coefficient schedule for GWO and WOA.",
                   ["Method", "Dataset", "Exp.", "Metric", "Median reference", "Median a0",
                    "Median paired difference", "p (raw)", "p (Holm)", "Unique evaluations reference / a0"], rows, True,
                   "a0: coefficient a reaches exactly 0 on the last evaluated update (GWO collapses to one point, "
                   "wasting the last 9 calls; WOA re-evaluates the best position); reference: a = 2 − t·(2/T), last "
                   "evaluated update a = 0.4. Paired difference oriented so that > 0 favours the reference schedule; "
                   "paired Wilcoxon test (zero_method = pratt), Holm correction over the 10 dataset × experiment cells "
                   f"per method. Tuning-versus-default conclusions that change when a0 replaces the reference variants: "
                   f"{int(rq1.changed.sum())} of {len(rq1)}.", 8.0))

    # S8 compute
    t = order_key(pd.read_csv(FINAL / "9_compute.csv"))
    comp = fs["compute"]
    rows = [[r.dataset, r.method, f3(r.median_s_per_call), f"{f3(r.median_unique)} ({int(r.min_unique)})",
             f3(r.total_cpu_h_clean), f3(r.total_cpu_h_raw), int(r.dropped_calls)] for r in t.itertuples()]
    rows.append(["All", "All", "–", "–", f3(comp["total_cpu_h_clean"]), f3(comp["total_cpu_h_raw"]),
                 int(comp["calls_dropped"])])
    tables.append(("Compute cost per dataset and optimizer.",
                   ["Dataset", "Method", "Median s per fitted call", "Median (min) unique evaluations per run",
                    "Worker-hours, sleep removed", "Worker-hours, raw", "Calls dropped (> 600 s)"], rows, False,
                   "Primary runs only (60 runs per dataset × method: 2 experiments × 30 seeds). Time per call measured "
                   "under 10 concurrent workers (Intel i7-10750H, 6 cores / 12 threads), one XGBoost thread per model. "
                   f"Calls longer than 600 s contain machine-sleep periods and were dropped ({comp['calls_dropped']} calls "
                   f"in {comp['runs_affected']} runs).", 8.0))

    # S9 headroom (exploratory)
    h = fs["headroom_EXPLORATORY"]
    t = pd.DataFrame(h["cells"])
    t["_d"] = t.dataset.map({d: i for i, d in enumerate(DS)})
    t = t.sort_values(["_d", "experiment"])
    rows = [[r.dataset, r.experiment, f3(r.default_val), f3(r.median_best_val_optimizers), f3(r.val_headroom_pct),
             f3(r.default_test), f3(r.median_test_optimizers), f3(r.test_improvement_pct)] for r in t.itertuples()]
    tables.append(("Validation headroom versus test gain (exploratory, not pre-specified).",
                   ["Dataset", "Exp.", "Default validation", "Median best validation (5 optimizers)",
                    "Validation headroom (%)", "Default test", "Median test (5 optimizers)", "Test gain (%)"], rows, False,
                   "Exploratory, not pre-specified. Headroom and test gain are oriented so that > 0 means the optimizers "
                   "are better than the default (RMSE: default − optimizers; ROC-AUC: optimizers − default) and are "
                   "expressed relative to the default's value. Spearman correlation over the 10 cells: "
                   f"ρ = {f3(h['spearman_rho_pct'])}, p = {fp(h['p_pct'])} (relative); ρ = {f3(h['spearman_rho_raw_units'])}, "
                   f"p = {fp(h['p_raw_units'])} (raw units); no multiplicity correction.", 8.0))

    # S10 curation (values from analysis/final/curation/curation_summary.md)
    detail = pd.read_csv(FINAL / "curation" / "curation_duplicates_detail.csv")
    cur = {  # dataset: (n, [train, val, test] duplicate groups (extra copies), label note)
        "ESOL": (1128, ["9 (9)", "2 (2)", "0"], "max |Δ logS| within groups: train 1.03, val 0.10"),
        "FreeSolv": (642, ["0", "0", "0"], "–"),
        "Lipophilicity": (4200, ["0", "0", "0"], "–"),
        "BBBP": (2039, ["62 (66)", "0", "1 (1)"], "11 train groups with conflicting classes"),
        "BACE": (1513, ["0", "0", "0"], "–"),
    }
    rows = []
    for ds in DS:
        n, within, note = cur[ds]
        cells = []
        for pair in ("train-valid", "train-test", "valid-test"):
            ex = detail[(detail.dataset == ds) & (detail.pair == pair) & detail.type.str.startswith("exact")]
            va = detail[(detail.dataset == ds) & (detail.pair == pair) & detail.type.str.startswith("stereo")]
            groups = va.inchikey.map(lambda k: k.split(" | ")[0][:14]).nunique()
            same = int((ex.y_a == ex.y_b).sum() + (va.y_a == va.y_b).sum())
            cells.append(f"{ex.inchikey.nunique()} / {groups} ({len(va)})" + (f"; labels {same}/{len(ex) + len(va)} agree"
                                                                               if len(ex) + len(va) else ""))
        rows.append([ds, n] + within + cells + [note])
    tables.append(("Data-curation check: duplicate molecules within and across the scaffold splits.",
                   ["Dataset", "Molecules", "Duplicate groups (extra copies) train", "val", "test",
                    "Cross-split train–val: exact / variant groups (pairs)", "train–test", "val–test",
                    "Label notes, within-split duplicates"], rows, True,
                   "Read-only check on the exact molecules and split used in the study (analysis/final/curation). "
                   "Duplicate: same standard InChIKey (RDKit 2026.03.6). Variant: same InChIKey first block "
                   "(connectivity) but different full key, i.e., different stereochemistry and/or protonation state; of "
                   "the 12 BACE variant pairs, 2 differ only in stereochemistry, 2 only in protonation and 8 in both. "
                   "Such pairs can cross the Bemis–Murcko scaffold split because scaffolds were computed from the SMILES "
                   "as written. Total: 1 exact cross-split duplicate (BBBP, train–val) and 11 variant groups (BACE); all "
                   "cross-split labels agree.", 7.5))
    return tables


def main() -> None:
    check_convergence_claim()
    fs = json.loads((FINAL / "final_stats.json").read_text(encoding="utf-8"))
    doc = Document()
    doc.core_properties.title, doc.core_properties.author = TITLE, AUTHOR
    doc.core_properties.last_modified_by = AUTHOR
    for name in ("Normal", "Caption"):
        st = doc.styles[name]
        st.font.name, st.font.size = "Times New Roman", Pt(11 if name == "Normal" else 10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    doc.styles["Caption"].font.color.rgb = None
    doc.styles["Caption"].font.italic = False
    s0 = doc.sections[0]
    s0.page_width, s0.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s0, side, MARGIN)
    fp_ = s0.footer.paragraphs[0]
    fp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp_.add_run("S")
    add_field(fp_, "PAGE", "1")

    # title page + contents
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(TITLE); r.bold = True; r.font.size = Pt(16)
    t.paragraph_format.space_before, t.paragraph_format.space_after = Pt(72), Pt(18)
    a = doc.add_paragraph(); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ra = a.add_run(AUTHOR); ra.font.size = Pt(13)
    a.paragraph_format.space_after = Pt(36)
    for heading, label in (("List of Figures", "Figure"), ("List of Tables", "Table")):
        h = doc.add_paragraph(); rh = h.add_run(heading); rh.bold = True; rh.font.size = Pt(13)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(12), Pt(6)
        add_field(doc.add_paragraph(), f'TOC \\h \\z \\c "{label}"', "Update fields to show the list.")

    # figures
    from PIL import Image
    for i, (fname, title, text) in enumerate(FIGURES, start=1):
        new_section(doc, landscape=False)
        with Image.open(FIGS / fname) as im:
            aspect = im.height / im.width
        width = min(FIG_MAX_W, FIG_MAX_H / aspect)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(FIGS / fname), width=Inches(width))
        caption(doc, "Figure", i, title)
        legend(doc, text)

    # tables
    for i, (title, header, rows, landscape, note, font_pt) in enumerate(build_tables(fs), start=1):
        new_section(doc, landscape=landscape)
        caption(doc, "Table", i, title)
        add_table(doc, header, rows, font_pt, total_in=9.5 if landscape else 7.0)
        legend(doc, note).paragraph_format.space_before = Pt(6)

    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT / "SI.docx")
    print("saved", OUT / "SI.docx", "| figures:", len(FIGURES), "| tables:", len(build_tables(fs)))


if __name__ == "__main__":
    main()
